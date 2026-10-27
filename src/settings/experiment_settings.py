"""
Experiment Settings Module

The experiment settings say which animals are recorded by each camera in
each recording round. They are imported from an Excel file (.xlsx) with
the columns:

    recording_round_number | camera | animals

- camera:  the camera nickname shown in the application (e.g. box1)
- animals: one or more animal IDs, separated by commas or semicolons
           (e.g. "M1, M2, M3, M4"). Several rows with the same round and
           camera are combined, so one animal per row also works.

Each video is named:  [YYYYMMDD]_[HHMMSS]_[camera]_[round]_[animal0-animal1-...].mp4
Experiment settings are optional; without them the name is [YYYYMMDD]_[HHMMSS]_[camera].mp4

Imported settings can be saved as a JSON file and loaded again later.
"""

import json
import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook


REQUIRED_COLUMNS = ('recording_round_number', 'camera', 'animals')

# Round numbers and camera names: letters, digits and "-" only ("_" separates name parts).
_NAME_PART_PATTERN = re.compile(r"^[A-Za-z0-9-]+$")
# Characters Windows does not allow in file names.
_FILENAME_FORBIDDEN = set('<>:"/\\|?*')


@dataclass
class ExperimentSettings:
    """Animals per (round, camera)."""
    source_file: str
    # One entry per (round, camera): {'round': '1', 'camera': 'box1', 'animals': ['M1', 'M2']}
    assignments: List[Dict[str, object]] = field(default_factory=list)

    def rounds(self) -> List[str]:
        """Round numbers in ascending order."""
        rounds = {a['round'] for a in self.assignments}
        return sorted(rounds, key=lambda r: (0, int(r), "") if r.isdigit() else (1, 0, r))

    def camera_names(self) -> List[str]:
        """Camera names (compared without case, like in animals_for)."""
        unique = {a['camera'].lower(): a['camera'] for a in self.assignments}
        return sorted(unique.values(), key=str.lower)

    def animals_for(self, round_id: str, camera_name: str) -> Optional[List[str]]:
        """Animals of a camera in a round (camera name compared without case), or None."""
        for a in self.assignments:
            if a['round'] == round_id and a['camera'].lower() == camera_name.lower():
                return a['animals']
        return None

    def save(self, json_path: str):
        data = {
            'saved_at': datetime.now().isoformat(timespec='seconds'),
            'source_file': self.source_file,
            'assignments': self.assignments,
        }
        with open(json_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, indent=2, ensure_ascii=False)

    @classmethod
    def load(cls, json_path: str) -> 'ExperimentSettings':
        with open(json_path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        return cls(source_file=data['source_file'], assignments=data['assignments'])


def video_name_stem(round_id: str, camera_name: str, animals: List[str]) -> str:
    """Name without date/time. Example: video_name_stem('3', 'box1', ['M1', 'M2']) -> 'box1_3_M1-M2'"""
    return f"{camera_name}_{round_id}_{'-'.join(animals)}"


def _cell_text(value) -> str:
    """Text of an Excel cell; whole numbers stored as 3.0 become '3'."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def import_experiment_excel(excel_path: str) -> Tuple[ExperimentSettings, List[str]]:
    """
    Read and check an experiment Excel file (first sheet).

    Returns:
        (settings, warnings) - warnings do not block the import.

    Raises:
        ValueError: listing every problem found, if the file cannot be used.
    """
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        rows = list(workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        workbook.close()

    # The header is the first non-empty row.
    while rows and all(_cell_text(v) == "" for v in rows[0]):
        rows.pop(0)
    if not rows:
        raise ValueError("The Excel file is empty.")
    header = [_cell_text(v).lower() for v in rows[0]]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise ValueError(f"Missing column(s): {', '.join(missing)}. "
                         f"Required columns: {', '.join(REQUIRED_COLUMNS)}. Found: {', '.join(h for h in header if h)}")
    col = {name: header.index(name) for name in REQUIRED_COLUMNS}

    errors = []
    warnings = []
    merged: Dict[Tuple[str, str], Dict[str, object]] = {}

    for excel_row_number, row in enumerate(rows[1:], start=2):
        values = {name: _cell_text(row[i]) if i < len(row) else "" for name, i in col.items()}
        if not any(values.values()):
            continue  # Empty line
        round_id, camera, animals_text = values['recording_round_number'], values['camera'], values['animals']

        if not round_id or not camera or not animals_text:
            errors.append(f"Row {excel_row_number}: empty cell (every row needs round, camera and animals).")
            continue
        if not _NAME_PART_PATTERN.match(round_id):
            errors.append(f"Row {excel_row_number}: round '{round_id}' may only contain letters, digits and '-'.")
        if not _NAME_PART_PATTERN.match(camera):
            errors.append(f"Row {excel_row_number}: camera '{camera}' may only contain letters, digits and '-'.")

        animals = [a.strip() for a in re.split(r"[,;\n]", animals_text) if a.strip()]
        for animal in animals:
            if any(ch in _FILENAME_FORBIDDEN for ch in animal):
                errors.append(f"Row {excel_row_number}: animal '{animal}' contains a character not allowed "
                              f"in file names ({' '.join(sorted(_FILENAME_FORBIDDEN))}).")
            elif '_' in animal or '-' in animal:
                warnings.append(f"Row {excel_row_number}: animal '{animal}' contains '_' or '-', which are also "
                                f"used as separators in the video name.")

        key = (round_id, camera.lower())
        if key in merged:
            merged[key]['animals'].extend(animals)
            warnings.append(f"Row {excel_row_number}: round {round_id} / camera {camera} appears in several rows; "
                            f"the animals are combined.")
        else:
            merged[key] = {'round': round_id, 'camera': camera, 'animals': animals}

    if errors:
        raise ValueError("\n".join(errors))
    if not merged:
        raise ValueError("The Excel file has no data rows.")
    return ExperimentSettings(source_file=str(Path(excel_path).resolve()), assignments=list(merged.values())), warnings
